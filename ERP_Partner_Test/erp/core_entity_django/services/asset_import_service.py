import csv
import io
import uuid
from decimal import Decimal
from datetime import datetime

from django.core.exceptions import ValidationError
from django.db import transaction

from openpyxl import load_workbook

from core_entity_django.models import (
    AssetCore,
    AssetImportBatch,
    AssetImportError,
    AssetAuditLog,
)

from core_entity_django.services.asset_core_services import (
    create_asset,
)


REQUIRED_COLUMNS = {
    "asset_name",
    "asset_type",
    "quantity",
    "unit_of_measure",
    "asset_class",
    "asset_category",
    "company_legal_entity",
    "purchase_cost",
    "currency",
}


def _parse_date(value):
    if value in (None, ""):
        return None

    if hasattr(value, "date"):
        return value.date()

    value = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(
                value,
                fmt,
            ).date()
        except ValueError:
            continue

    raise ValueError(
        f"Invalid date: {value}"
    )


def _load_rows(uploaded_file):

    filename = uploaded_file.name.lower()

    if filename.endswith(".csv"):

        content = uploaded_file.read().decode(
            "utf-8-sig"
        )

        reader = csv.DictReader(
            io.StringIO(content)
        )

        return list(reader), "CSV"

    if filename.endswith(".xlsx"):

        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            return [], "XLSX"

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in rows[0]
        ]

        result = []

        for values in rows[1:]:
            result.append(
                dict(
                    zip(
                        headers,
                        values,
                    )
                )
            )

        return result, "XLSX"

    raise ValidationError(
        "File format must be Excel (.xlsx) or CSV."
    )


def _normalize(row):

    return {
        str(key).strip(): value
        for key, value in row.items()
    }


def _row_to_asset_data(row):

    row = _normalize(row)

    missing = [
        column
        for column in REQUIRED_COLUMNS
        if row.get(column) in (None, "")
    ]

    if missing:
        raise ValidationError(
            f"Required fields are missing: {', '.join(sorted(missing))}"
        )

    try:
        quantity = Decimal(
            str(row["quantity"])
        )

        purchase_cost = Decimal(
            str(row["purchase_cost"])
        )

        other_cost = Decimal(
            str(
                row.get(
                    "other_directly_attributable_cost",
                    0,
                )
            )
        )

    except Exception:
        raise ValidationError(
            "Quantity and cost fields must be numeric."
        )

    if quantity <= 0:
        raise ValidationError(
            "Quantity must be greater than zero."
        )

    if purchase_cost <= 0:
        raise ValidationError(
            "Purchase cost must be greater than zero."
        )

    entry_source = AssetCore.SourceType.IMPORT

    return {
        "draft_date": _parse_date(
            row.get("draft_date")
        ),
        "entry_source": entry_source,
        "remarks": row.get("remarks", "") or "",
        "basic_info": {
            "asset_name": row["asset_name"],
            "asset_description": row.get(
                "asset_description",
                "",
            ) or "",
            "brand": row.get("brand", "") or "",
            "model": row.get("model", "") or "",
            "serial_number": row.get(
                "serial_number",
                "",
            ) or "",
            "manufacturer": row.get(
                "manufacturer",
                "",
            ) or "",
            "quantity": quantity,
            "unit_of_measure": row["unit_of_measure"],
            "asset_type": row["asset_type"],
            "component_flag": bool(
                row.get("component_flag", False)
            ),
        },
        "classification": {
            "asset_class": row["asset_class"],
            "asset_category": row["asset_category"],
            "asset_subcategory": row.get(
                "asset_subcategory",
                "",
            ) or "",
            "asset_group": row.get(
                "asset_group",
                "",
            ) or "",
            "depreciation_policy_group": row.get(
                "depreciation_policy_group",
                "",
            ) or "",
            "useful_life": row.get(
                "useful_life"
            ),
            "residual_value": row.get(
                "residual_value"
            ),
            "depreciation_method": row.get(
                "depreciation_method",
                "",
            ) or "",
            "capitalization_threshold_check": bool(
                row.get(
                    "capitalization_threshold_check",
                    False,
                )
            ),
            "tax_asset_category": row.get(
                "tax_asset_category",
                "",
            ) or "",
        },
        "location_info": {
            "company_legal_entity": row[
                "company_legal_entity"
            ],
            "branch": row.get(
                "branch",
                "",
            ) or "",
            "department": row.get(
                "department",
                "",
            ) or "",
            "location": row.get(
                "location",
                "",
            ) or "",
            "building_floor_room": row.get(
                "building_floor_room",
                "",
            ) or "",
            "rack_space_area": row.get(
                "rack_space_area",
                "",
            ) or "",
            "current_physical_status": row.get(
                "current_physical_status",
                "",
            ) or "",
        },
        "financial_info": {
            "acquisition_date": _parse_date(
                row.get("acquisition_date")
            ),
            "invoice_date": _parse_date(
                row.get("invoice_date")
            ),
            "invoice_number": row.get(
                "invoice_number",
                "",
            ) or "",
            "purchase_order_number": row.get(
                "purchase_order_number",
                "",
            ) or "",
            "purchase_cost": purchase_cost,
            "other_directly_attributable_cost": other_cost,
            "currency": row["currency"],
            "exchange_rate": row.get(
                "exchange_rate"
            ),
            "capitalization_date": _parse_date(
                row.get("capitalization_date")
            ),
            "gl_asset_account": row.get(
                "gl_asset_account",
                "",
            ) or "",
            "accumulated_depreciation_account": row.get(
                "accumulated_depreciation_account",
                "",
            ) or "",
            "depreciation_expense_account": row.get(
                "depreciation_expense_account",
                "",
            ) or "",
            "cost_center": row.get(
                "cost_center",
                "",
            ) or "",
            "funding_source": row.get(
                "funding_source",
                "",
            ) or "",
        },
        "supplier_info": {
            "supplier_name": row.get(
                "supplier_name",
                "",
            ) or "",
            "supplier_code": row.get(
                "supplier_code",
                "",
            ) or "",
            "procurement_method": row.get(
                "procurement_method",
                "",
            ) or "",
            "contract_number": row.get(
                "contract_number",
                "",
            ) or "",
            "delivery_note_number": row.get(
                "delivery_note_number",
                "",
            ) or "",
            "warranty_start_date": _parse_date(
                row.get("warranty_start_date")
            ),
            "warranty_end_date": _parse_date(
                row.get("warranty_end_date")
            ),
            "warranty_provider": row.get(
                "warranty_provider",
                "",
            ) or "",
        },
        "responsibility": {
            "responsible_department": row.get(
                "responsible_department",
                "",
            ) or "",
            "custodian_employee": row.get(
                "custodian_employee",
                "",
            ) or "",
            "asset_controller": row.get(
                "asset_controller",
                "",
            ) or "",
            "user_department": row.get(
                "user_department",
                "",
            ) or "",
            "assigned_user": row.get(
                "assigned_user",
                "",
            ) or "",
        },
    }


@transaction.atomic
def import_assets(
    *,
    user,
    uploaded_file,
    template_version,
    source_ip=None,
    device="",
):

    rows, file_type = _load_rows(
        uploaded_file
    )

    batch = AssetImportBatch.objects.create(
        batch_number=(
            f"IMP-{uuid.uuid4().hex[:10].upper()}"
        ),
        file_name=uploaded_file.name,
        file_type=file_type,
        template_version=template_version,
        uploaded_by=user,
        total_rows=len(rows),
    )

    created_assets = []
    errors = []

    for index, row in enumerate(rows, start=2):

        try:
            data = _row_to_asset_data(row)

            asset = create_asset(
                user=user,
                validated_data=data,
                source_ip=source_ip,
                device=device,
                batch_number=batch.batch_number,
            )

            created_assets.append(asset)

        except Exception as exc:

            message = str(exc)

            AssetImportError.objects.create(
                batch=batch,
                row_number=index,
                field_name="",
                error_code="IMPORT_ROW_ERROR",
                error_message=message,
                error_type=(
                    AssetImportError.ErrorType.BLOCKING
                ),
                suggested_correction=(
                    "Correct the row and import again."
                ),
            )

            errors.append(
                {
                    "row_number": index,
                    "error": message,
                }
            )

    batch.valid_rows = len(created_assets)
    batch.error_rows = len(errors)

    if errors:
        batch.status = (
            AssetImportBatch.Status.FAILED
            if not created_assets
            else AssetImportBatch.Status.VALIDATED
        )
    else:
        batch.status = (
            AssetImportBatch.Status.COMPLETED
        )

    batch.save()

    AssetAuditLog.objects.create(
        asset=None,
        action_type=AssetAuditLog.ActionType.IMPORT,
        user=user,
        new_value={
            "batch_number": batch.batch_number,
            "total_rows": batch.total_rows,
            "valid_rows": batch.valid_rows,
            "error_rows": batch.error_rows,
        },
        batch_number=batch.batch_number,
        source_ip=source_ip,
        device=device,
    )

    return {
        "batch": batch,
        "created_assets": created_assets,
        "errors": errors,
    }